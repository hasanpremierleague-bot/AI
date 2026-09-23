from pathlib import Path
import json, base64
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
R = json.loads((ROOT/'pipeline_results.json').read_text(encoding='utf-8'))
N = json.loads((ROOT/'student_academic_performance_pipeline.ipynb').read_text(encoding='utf-8'))
ART = ROOT/'revision_work/report_figures'
ART.mkdir(exist_ok=True)
figures = {}
for i, cell in enumerate(N['cells']):
    for j, out in enumerate(cell.get('outputs', [])):
        png = out.get('data', {}).get('image/png')
        if png:
            path = ART/f'cell_{i}_{j}.png'
            path.write_bytes(base64.b64decode(png))
            figures.setdefault(i, []).append(path)

d = Document()
sec = d.sections[0]
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
style = d.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(7)
style.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 25), ('Heading 1', 17), ('Heading 2', 12)]:
    d.styles[name].font.name = 'Calibri'; d.styles[name].font.size = Pt(size)
    d.styles[name].font.color.rgb = RGBColor.from_string('000000' if name=='Title' else '183B56')
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer = sec.footer.paragraphs[0]; footer.alignment = 2
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text): return d.add_paragraph(text)
def h(text): return d.add_heading(text, 2)
def page(title): d.add_page_break(); d.add_heading(title, 1)
def table(headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    t.autofit=False
    for i,x in enumerate(headers): t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    for k,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit');pr.append(no)
        for i,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                for run in para.runs:
                    run.font.size=Pt(9)
                    if k==0: run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'183B56' if k==0 else ('EEF3F7' if k%2 else 'FFFFFF'))
            c._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p('')
    return t
def pic(cell, caption, idx=0, width=6.3):
    para=d.add_paragraph();para.alignment=1
    para.add_run().add_picture(str(figures[cell][idx]),width=Inches(width))
    para.paragraph_format.keep_with_next=True
    cap=p(caption);cap.style='Caption'
def fmt(x):return f'{x:.4f}'
def params(x):return ', '.join(f'{k.replace("model__", "")}={v}' for k,v in x.items())

d.add_heading('Academic performance analysis and outcome prediction',0)
p('IT7103 Advanced AI Applications | Project report')
p('Bahrain Polytechnic | Stream 05 | Group 1')
p('Assessment deadline: 15 December 2026, 11:55 PM')
table(['Student ID','Student name'],[
['202303596','Hasan Marhoon (group leader)'],['202304831','Sayed Ali Almusawi'],
['202304829','Hussain Ali'],['202305013','Sayed Sadeq'],['202304419','Ali Yusuf']], [1.6,4.9])
h('Overview')
p('This project analyzes the supplied student-course dataset, predicts examination scores and letter grades, and compares individual time-series forecasts. The revised evaluation preserves observed targets, uses chronological holdouts, and fits preprocessing separately within each tuning fold. Results support a limited retrospective prediction experiment; they do not establish causal effects or a validated early-warning service.')
p('Companion notebook: student_academic_performance_pipeline.ipynb. All result tables and figures in this report are generated from its executed outputs. The assessment brief and supplied CSV are the task and data sources [1, 2].')
h('Report organization')
p('Problem and data; preparation and validation; exploratory analysis; regression; classification; forecasting; optimization; interpretation and conclusions. These sections map to the brief without assuming an unsupported allocation of marks.')

page('Problem objectives and data preparation')
p('The practical problem is identifying patterns associated with academic performance in a large set of student-course records. The objectives are to predict final_exam, classify grade_letter, and forecast one student’s semester means while comparing individual models with ensembles. A useful analysis must distinguish model accuracy from evidence that an intervention improves outcomes.')
s=R['dataset']
table(['Dataset measure','Observed value'], [[k.replace('_',' ').capitalize(), f'{s[k]:,}'] for k in ['raw_rows','clean_rows','students','raw_observed_exams','clean_observed_exams','missing_exams_clean']], [4.4,2.1])
p('Each row represents a course enrollment. Predictors include GPA, study hours, attendance, homework, course credits and categorical student/course attributes. Scholarship and internship are binary indicators; extracurricular is a count, treated categorically in this implementation. No claim is made that the data are actual Bahrain Polytechnic records.')
h('Cleaning choices')
p(f'Remove 780 rows with ERR grades or numeric grades outside 0-100. Among retained rows, clip {s["bad_gpa_clean"]:,} invalid GPA values to 0-4 and cap {s["study_hours_capped"]:,} study-hour values at 60. These are explicit domain assumptions, not proof of the original values. An alternative is to set invalid GPA entries missing; sensitivity to these choices remains future work. The duplicate audit found no exact duplicate rows.')
p('Missing exam targets are never imputed. Regression uses observed exams only; classification can retain rows with valid letter grades even if the exam is missing. Numeric predictors use training medians, categorical predictors use the training mode, numeric scaling uses training statistics, and categorical encoding tolerates previously unseen values.')

page('Temporal features and validation design')
p('Semesters are mapped chronologically from 2020Spring (1) to 2023Fall (8). A complete student-by-semester grid preserves gaps as missing. Per-student means use observed scores. Lag features shift means by one term; two-semester rolling means use only t-2 and t-1. Missing histories are not backfilled from the future.')
p('Added predictors include prior exam, homework and attendance averages, rolling homework and attendance means, study hours divided by credits, and a 0.4 attendance + 0.6 homework composite. Current final_exam, numeric_grade and grade_letter are excluded from predictors.')
table(['Task','Training rows, semesters 1-6','Test rows, semesters 7-8'],[
['Regression',f'{s["reg_train_rows"]:,}',f'{s["reg_test_rows"]:,}'],
['Classification',f'{s["train_rows"]:,}',f'{s["test_rows"]:,}']], [1.5,2.5,2.5])
h('Information available at prediction time')
p('The main evaluation is sequential by semester: semester-8 features can use observed semester-7 history. Both test semesters are not forecast simultaneously at the end of semester 6. The trained model remains fixed. Current-semester homework and attendance are aggregate measures without within-term timestamps, so early-week availability has not been validated. This protocol concerns recurring students, not an independent unseen-student population.')
h('Whole-semester cross-validation')
table(['Training semesters','Validation semester'],[['1-3','4'],['1-4','5'],['1-5','6']], [4,2.5])
p('Regression and classification use these expanding folds, with every training semester earlier than its validation semester. Each GridSearchCV estimator is a complete preprocessing/model pipeline; imputers, scalers and encoders are refitted inside each fold [3]. Parameters are selected using only semesters 1-6 and then refitted on that training period.')
p('The historical holdout has been examined during project revision. It therefore serves as a documented benchmark, not pristine external evidence. A later cohort is needed before deployment claims.')

page('Exploratory associations and cohort trends')
c=R['correlations']; means=R['cohort']['final_exam']
table(['Association with observed final exam','Pearson correlation'],[[x,fmt(c[k])] for x,k in [('Homework','homework_avg'),('Attendance','attendance_rate'),('Previous semester exam','final_exam_prev_semester')]], [4.5,2])
p(f'Semester mean exam scores range from {min(means.values()):.2f} to {max(means.values()):.2f}. Homework has a substantial positive association, while attendance and previous-semester exam averages have weak linear associations in these data. Population means do not describe every student’s trajectory. Correlation does not establish causation.')
pic(12,'Figure 1. Observed-value correlations, including engineered features.',width=5.7)
p('EDA summarizes the full historical dataset retrospectively. The search spaces are fixed independently of this holdout analysis. Missing values are omitted pairwise when computing correlations.')

page('Regression results and feature comparison')
p('Ridge provides a regularized linear baseline; a decision tree captures nonlinear partitions. Random Forest averages trees, LightGBM and XGBoost boost trees, and the voting model combines RF, LightGBM and XGBoost with weights 1:2:2. A training-mean predictor supplies a minimal reference. RMSE and MAE are in exam-score points; R2 measures explained variation and MAPE is a percentage.')
reg=dict(R['regression']);reg['Tuned LightGBM']=R['tuning_regression']['Tuned LightGBM']
table(['Model','RMSE','MAE','R2','MAPE %'],[[name]+[fmt(v[k]) for k in ['RMSE','MAE','R2','MAPE (%)']] for name,v in reg.items()],[2.5,1,1,1,1])
winner=min(reg,key=lambda k:reg[k]['RMSE'])
p(f'The lowest observed test RMSE is {reg[winner]["RMSE"]:.4f} for {winner}. Small differences between models should not be treated as established superiority without repeated external validation. Greater algorithmic complexity does not guarantee a better result.')
h('Added feature ablation')
table(['Model','Base RMSE','Added features RMSE','Base minus added'],[[name]+[fmt(v[k]) for k in ['Raw Features RMSE','Engineered Features RMSE','Delta (Lift)']] for name,v in R['ablation'].items()],[2.1,1.4,1.6,1.4])
p('The added set contains temporal features and interactions together. Its very small differences do not demonstrate a meaningful lift and do not isolate the effect of lag features alone. The ablation LightGBM uses 100 trees, while the main baseline uses 150; comparisons within each ablation pair use identical settings.')

page('Classification results and alert tradeoffs')
p('The ten classes use one explicit mapping throughout: A, A-, B+, B, B-, C+, C, C-, D, F. Each unweighted/balanced pair has identical baseline settings. Balanced weights are calculated from the training labels. Macro metrics average classes equally; F-specific metrics quantify failure detection.')
table(['Model','Accuracy','Macro precision','Macro recall','Macro F1'],[[name]+[fmt(v[k]) for k in ['Overall Acc','Macro Precision','Macro Recall','Macro F1']] for name,v in R['classification'].items()],[2.5,1,1,1,1])
table(['Model','F recall','F precision'],[[name,fmt(v['Class F Recall']),fmt(v['Class F Precision'])] for name,v in R['classification'].items() if 'Balanced' in name or name=='Random Forest (Unweighted)'],[3.7,1.4,1.4])
rf=R['classification']['Random Forest (Balanced)']
p(f'Balanced Random Forest detects {rf["Class F Recall"]:.1%} of observed F grades, with precision {rf["Class F Precision"]:.1%}. Thus approximately {1-rf["Class F Precision"]:.1%} of its F alerts are false positives. This tradeoff must be weighed against advising capacity. The unweighted RF comparison supports its own recall improvement; other unweighted models do detect some failures.')

page('Classification and regression diagnostics')
pic(19,'Figure 2. Balanced Random Forest confusion matrix with the explicit grade order.',width=5.0)
pic(16,'Figure 3. Ridge residuals and residual distribution on observed test targets.',width=6.3)
p('The confusion matrix shows which adjacent grades are confused. Residual plots help identify bias and changing error spread; visual appearance alone does not prove normality, independence or equal variance. The Ridge scatter displays the first 3,000 test predictions for legibility.')

page('Individual student forecasting')
p('Student 2 has eight semesters and 33 enrolled courses. Semester means use only observed exams, and course counts need not equal observed-exam counts. The first six semester means train SES and ARIMA(1,0,0); the final two are forecast without updating from the first held-out result. Persistence repeats the final training value [5].')
table(['Semester','Courses','Observed exams','Exam mean'],[[v['semester'],v['course_count'],v['observed_exam_count'],f'{v["final_exam"]:.2f}'] for v in R['student_history']],[2.2,1.2,1.6,1.5])
table(['Forecast','RMSE','MAE','MAPE %'],[[name]+[fmt(v[k]) for k in ['RMSE','MAE','MAPE (%)']] for name,v in R['forecast'].items()],[3.2,1.1,1.1,1.1])
pic(21,'Figure 4. Fixed-origin forecasts compared with observed semester means.',width=6)
p('Only two test observations are available. The lowest observed error is descriptive evidence for this case, not a general result about forecasting methods. Six training points leave substantial parameter uncertainty; rolling-origin evaluation over longer histories and more students is needed.')

page('Optimization results')
p('Regression searches 8 configurations: 100/200 trees, learning rate 0.03/0.08, and maximum depth 4/6. Classification searches 4 configurations: 100 trees, depth 10/14, and minimum leaf size 1/5. Each configuration uses the three whole-semester folds. Search preprocessing is learned inside each fold; both selected models are refitted on all eligible training records.')
p('LightGBM subsample is not searched because subsample_freq defaults to zero, which disables row subsampling [4]. Searches raise fitting errors rather than silently accepting failed folds.')
h('Regression tuning')
p('Selected parameters: '+params(R['reg_best_params'])+f'. Mean validation RMSE: {R["reg_cv_rmse"]:.4f}.')
table(['Metric','Baseline LightGBM','Tuned LightGBM'],[[k,fmt(R['tuning_regression']['Baseline LightGBM'][k]),fmt(R['tuning_regression']['Tuned LightGBM'][k])] for k in ['RMSE','MAE','R2','MAPE (%)']],[2.5,2,2])
h('Classification tuning')
p('Selected parameters: '+params(R['clf_best_params'])+f'. Mean validation macro F1: {R["clf_cv_f1"]:.4f}.')
table(['Metric','Baseline balanced RF','Tuned balanced RF'],[[k,fmt(R['tuning_classification']['Baseline balanced RF'][k]),fmt(R['tuning_classification']['Tuned balanced RF'][k])] for k in ['Overall Acc','Macro Precision','Macro Recall','Macro F1','Class F Recall','Class F Precision']],[2.5,2,2])
delta=R['tuning_classification']['Tuned balanced RF']['Macro F1']-R['tuning_classification']['Baseline balanced RF']['Macro F1']
p(f'The tuned classifier changes test macro F1 by {delta:+.4f}. Validation selects parameters; the holdout measures the resulting model. Tuning may improve one metric while worsening another. These bounded searches are not exhaustive and should not be presented as finding a global optimum.')

page('Interpretation conclusions and reproducibility')
top=R['importance'][0]
p(f'The largest fitted LightGBM gain share is {top["Gain_share"]:.1%} for {top["Feature"]}. This describes the trained model, not a causal effect on student achievement. Correlated features can exchange importance. Split frequency counts branch decisions, whereas gain accumulates loss reduction [4].')
pic(25,'Figure 5. LightGBM gain and split importance from the selected regression pipeline.',width=6.3)
h('Conclusions and next improvements')
p('The project implements the requested cleaning, temporal features, regression, classification, forecasting, ensembles and optimization. Baseline comparisons constrain the interpretation: added features show little lift, class weighting increases failure recall at the cost of false alerts, and very short student histories limit forecasting evidence. No grade, intervention benefit or deployment readiness follows from completing these analyses.')
p('Next work should obtain later unseen records, validate the prediction timestamp, compare clipping policies, evaluate calibrated probabilities and subgroup errors, and assess advising workload. Fairness cannot be inferred from low feature importance. Student identifiers should be access-controlled, and any practical use should support human review.')
h('Reproducibility')
p('Upload Project1-EducationDataset.csv in Colab and run the notebook from the first cell. Setup checks modeling-library versions, seeds are fixed at 42, and parallel jobs are bounded. The report is generated from the notebook’s exported JSON and embedded figures. The notebook was executed locally in a fresh kernel; actual hosted Colab execution is not claimed.')
p('Recorded versions: '+', '.join(f'{k} {v}' for k,v in R['versions'].items())+'.')

page('Sources and requirement mapping')
table(['Brief objective','Report coverage'],[
['Cleaning and outliers','Problem objectives and data preparation'],
['Preprocessing and temporal engineering','Temporal features and validation design'],
['Correlation and time-series analysis','Exploratory associations; individual forecasting'],
['Regression and classification','Results and diagnostics sections'],
['Ensembles and comparison','Regression and classification baseline comparisons'],
['Hyperparameter optimization','Optimization results for both tasks'],
['Interpretation and limitations','Interpretation conclusions and reproducibility']], [2.5,4])
p('[1] IT7103 Advanced AI Applications, supplied project brief, Project-1-EDUCATION (1).pdf. Assignment requirements and dataset field descriptions.')
p('[2] Supplied Project1-EducationDataset.csv. Primary source of all counts, descriptive statistics, model inputs and observed targets.')
p('[3] scikit-learn documentation, Pipelines and composite estimators. https://scikit-learn.org/stable/modules/compose.html (accessed 23 September 2026).')
p('[4] LightGBM documentation, LGBMRegressor API. https://lightgbm.readthedocs.io/en/latest/pythonapi/lightgbm.LGBMRegressor.html (accessed 23 September 2026).')
p('[5] R. J. Hyndman and G. Athanasopoulos, Forecasting: Principles and Practice, third edition, section 5.2, Some simple forecasting methods. https://otexts.com/fpp3/simple-methods.html (accessed 23 September 2026).')
d.save(ROOT/'Project_Report_IT7103.docx')
print('Report rebuilt from executed notebook results')
